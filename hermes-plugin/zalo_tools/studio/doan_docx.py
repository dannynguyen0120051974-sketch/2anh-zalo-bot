"""Bộ sinh CỐ ĐỊNH văn bản Đoàn trường (profile ``doan`` của skill ``soan-van-ban-doan``).

AI chỉ viết JSON (``check_doan_json``); mọi thể thức do mã này đặt, theo mẫu ưu tiên đã chốt của skill
(``Mau-Ke-hoach-Doan-hanh-chinh-ket-hop.docx`` + ``references/the-thuc-van-ban-doan.md``):

- A4, lề trên/dưới 20 mm, trái 30 mm, phải 20 mm; Times New Roman, chữ đen; số trang giữa đầu trang từ trang 2.
- Bảng hai cột ẩn viền 7/9 cm: trái = đơn vị cấp trên + "BAN CHẤP HÀNH ĐOÀN …" + ``---***---``; phải = quốc hiệu,
  tiêu ngữ, ``____________________``. Hàng 2: ``Số: N/KH-ĐTN`` | địa danh, ngày (nghiêng, cách trên 6 pt).
- Tên loại 16 pt đậm giữa; trích yếu 14 pt đậm giữa. Thân 14 pt, căn đều, giãn 1,15, trước/sau 3 pt, thụt
  dòng đầu 1,25 cm, lề trái 0; gạch đầu dòng gõ tay, không danh sách tự động, không tab.
- Khối ký: "Nơi nhận:" đậm nghiêng 12 pt + mục 11 pt | quyền hạn đậm 12 pt, chức vụ ("Bí thư") KHÔNG đậm, 4 dòng
  trống, họ tên đậm.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

TYPES = {  # loai -> (tên loại in trên văn bản, mã trong số ký hiệu)
    "ke_hoach": ("KẾ HOẠCH", "KH"),
    "thong_bao": ("THÔNG BÁO", "TB"),
    "cong_van": ("", "CV"),
    "bao_cao": ("BÁO CÁO", "BC"),
    "trieu_tap": ("GIẤY TRIỆU TẬP", "TrT"),
    "huong_dan": ("HƯỚNG DẪN", "HD"),
    "quyet_dinh": ("QUYẾT ĐỊNH", "QĐ"),
}
FONT = "Times New Roman"


def _run(paragraph, text: str, *, size: float, bold: bool = False, italic: bool = False):
    from docx.shared import Pt, RGBColor

    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if fonts is None:
        from docx.oxml import OxmlElement
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        from docx.oxml.ns import qn
        fonts.set(qn(attr), FONT)
    return run


def _para(cell_or_doc, text: str = "", *, size: float = 14, bold=False, italic=False, align="center",
          before: float = 0, after: float = 0, first: float = 0, line: float = 1.0):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    p = cell_or_doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT,
                     "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.left_indent = Cm(0)
    fmt.first_line_indent = Cm(first)
    fmt.line_spacing = line
    if text:
        _run(p, text, size=size, bold=bold, italic=italic)
    return p


def _hide_borders(table) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def _cell(table, row: int, col: int, width_cm: float):
    from docx.shared import Cm

    cell = table.cell(row, col)
    cell.width = Cm(width_cm)
    first = cell.paragraphs[0]
    first._element.getparent().remove(first._element)  # ô bắt đầu rỗng, mọi đoạn do _para thêm
    return cell


def _page_number_header(section) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    section.different_first_page_header_footer = True
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = _run(p, "", size=13)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._element.append(el)


def number_line(data: Dict[str, Any]) -> str:
    code = TYPES[data["loai"]][1]
    so = str(data.get("so") or "").strip()
    return f"Số: {so}/{code}-ĐTN" if so else f"Số:      /{code}-ĐTN"


def build(data: Dict[str, Any], out_path: Path) -> Path:
    """``data`` đã qua ``validate.check_doan_json``."""
    from docx import Document
    from docx.shared import Cm

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.0), Cm(2.0)
    section.left_margin, section.right_margin = Cm(3.0), Cm(2.0)
    _page_number_header(section)

    loai = data["loai"]
    ten_loai = TYPES[loai][0]
    head = doc.add_table(rows=2, cols=2)
    _hide_borders(head)
    left, right = _cell(head, 0, 0, 7.0), _cell(head, 0, 1, 9.0)
    _para(left, data["don_vi_cap_tren"], size=11.5, bold=True)
    _para(left, data["don_vi"], size=11.5, bold=True)
    _para(left, "---***---", size=11.5, bold=True)
    _para(right, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", size=11.5, bold=True)
    _para(right, "Độc lập - Tự do - Hạnh phúc", size=13, bold=True)
    _para(right, "____________________", size=12, bold=True)
    left2, right2 = _cell(head, 1, 0, 7.0), _cell(head, 1, 1, 9.0)
    _para(left2, number_line(data), size=13, before=6)
    if loai == "cong_van":
        _para(left2, f"V/v {data['trich_yeu']}", size=12)
    _para(right2, f"{data['dia_danh']}, ngày {data['ngay']} tháng {data['thang']} năm {data['nam']}",
          size=13, italic=True, before=6)

    if ten_loai:
        _para(doc, ten_loai, size=16, bold=True, before=12)
        _para(doc, data["trich_yeu"], size=14, bold=True, after=6)
    if data.get("kinh_gui"):
        names = data["kinh_gui"]
        if len(names) == 1:
            _para(doc, f"Kính gửi: {names[0]}", size=14, before=6, after=6)
        else:
            _para(doc, "Kính gửi:", size=14, before=6)
            for name in names:
                _para(doc, f"- {name}", size=14, align="left", first=3.5)

    for block in data["noi_dung"]:
        if "bang" in block:
            rows: List[List[str]] = block["bang"]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, text in enumerate(row):
                    cell = table.cell(r, c)
                    cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
                    _para(cell, text, size=13, bold=(r == 0), align="center" if r == 0 else "left")
            _para(doc, "", size=6)
            continue
        text = block.get("muc") or block.get("doan")
        _para(doc, text, size=14, bold="muc" in block, align="justify", before=3, after=3, first=1.25, line=1.15)
    if data.get("ket"):
        _para(doc, data["ket"], size=14, align="justify", before=3, after=3, first=1.25, line=1.15)

    sign = doc.add_table(rows=1, cols=2)
    _hide_borders(sign)
    nl, ky = _cell(sign, 0, 0, 7.0), _cell(sign, 0, 1, 9.0)
    _para(nl, "Nơi nhận:", size=12, bold=True, italic=True, align="left", before=12)
    for item in data["noi_nhan"]:
        _para(nl, f"- {item}", size=11, align="left")
    _para(ky, data["quyen_han"], size=12, bold=True, before=12)
    _para(ky, data["chuc_vu"], size=12)
    for _ in range(4):
        _para(ky, "", size=12)
    if data.get("nguoi_ky"):
        _para(ky, data["nguoi_ky"], size=13, bold=True)
    doc.core_properties.title = data["trich_yeu"][:200]
    doc.core_properties.author = data["don_vi"][:100]
    doc.save(str(out_path))
    return out_path
